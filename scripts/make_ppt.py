# -*- coding: utf-8 -*-
"""生成 MonadLens 演示 PPT（深色科技风，16:9）。

依赖：python-pptx
运行：python scripts/make_ppt.py
输出：MonadLens演示.pptx（位于仓库根目录）
约定：helper 函数参数统一用 sld；页面函数内用 s = slide_new()。
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------- 配色 ----------------
BG       = RGBColor(0x12, 0x10, 0x26)
BG_CARD  = RGBColor(0x22, 0x1E, 0x42)
PURPLE   = RGBColor(0x8B, 0x5C, 0xF6)
CYAN     = RGBColor(0x22, 0xD3, 0xEE)
AMBER    = RGBColor(0xF5, 0x9E, 0x0B)
DANGER   = RGBColor(0xEF, 0x44, 0x44)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
TEXT     = RGBColor(0xE9, 0xE9, 0xF2)
MUTED    = RGBColor(0x9C, 0xA3, 0xB8)
FONT     = "Microsoft YaHei"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide_new():
    return prs.slides.add_slide(BLANK)


def bg(sld, color=BG):
    r = sld.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background(); r.shadow.inherit = False
    return r


def text(sld, l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = sld.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.04); tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02); tf.margin_bottom = Inches(0.02)
    first = True
    for ln in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = ln.get("align", align)
        if "space_after" in ln: p.space_after = Pt(ln["space_after"])
        if "space_before" in ln: p.space_before = Pt(ln["space_before"])
        if "line_spacing" in ln: p.line_spacing = ln["line_spacing"]
        run = p.add_run(); run.text = ln["text"]
        f = run.font
        f.size = Pt(ln.get("size", 18)); f.color.rgb = ln.get("color", TEXT)
        f.bold = ln.get("bold", False); f.name = FONT
    return tb


def card(sld, l, t, w, h, fill=BG_CARD, line=None, radius=0.1):
    shp = sld.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(1.25)
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    shp.shadow.inherit = False
    return shp


def accent_bar(sld, l, t, h, color):
    bar = sld.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, Inches(0.09), h)
    bar.fill.solid(); bar.fill.fore_color.rgb = color
    bar.line.fill.background(); bar.shadow.inherit = False
    return bar


def header(sld, kicker, title):
    text(sld, Inches(0.6), Inches(0.42), Inches(12), Inches(0.36),
         [{"text": kicker, "size": 13, "color": PURPLE, "bold": True}])
    text(sld, Inches(0.6), Inches(0.78), Inches(12.1), Inches(0.85),
         [{"text": title, "size": 30, "color": WHITE, "bold": True}])


def footer(sld, n):
    text(sld, Inches(0.6), Inches(7.02), Inches(8), Inches(0.3),
         [{"text": "MonadLens · 看清网络，看清后果", "size": 10, "color": MUTED}])
    text(sld, Inches(11.9), Inches(7.02), Inches(0.9), Inches(0.3),
         [{"text": str(n), "size": 10, "color": MUTED, "align": PP_ALIGN.RIGHT}])


def bullets(items, size=15, color=TEXT, gap=7):
    out = []
    for it in items:
        if isinstance(it, tuple):
            txt, lvl = it
        else:
            txt, lvl = it, 0
        out.append({"text": ("•  " if lvl == 0 else "    –  ") + txt,
                    "size": size, "color": color, "space_after": gap})
    return out


def panel(sld, l, t, w, h, title, body, accent=PURPLE,
          title_size=18, body_size=14, fill=BG_CARD):
    card(sld, l, t, w, h, fill=fill)
    accent_bar(sld, l, t, h, accent)
    text(sld, l + Inches(0.26), t + Inches(0.14), w - Inches(0.42), Inches(0.5),
         [{"text": title, "size": title_size, "color": WHITE, "bold": True}])
    text(sld, l + Inches(0.26), t + Inches(0.72), w - Inches(0.5), h - Inches(0.9),
         body, anchor=MSO_ANCHOR.TOP)


# ================= 第 1 页：封面 =================
def page_cover():
    s = slide_new(); bg(s)
    card(s, 0, 0, SW, Inches(0.2), fill=PURPLE, radius=0)
    text(s, Inches(0.9), Inches(2.05), Inches(11.5), Inches(1.3),
         [{"text": "MonadLens", "size": 66, "color": WHITE, "bold": True,
           "align": PP_ALIGN.CENTER}])
    text(s, Inches(0.9), Inches(3.35), Inches(11.5), Inches(0.7),
         [{"text": "看清网络，看清后果", "size": 30, "color": PURPLE, "bold": True,
           "align": PP_ALIGN.CENTER}])
    text(s, Inches(0.9), Inches(4.35), Inches(11.5), Inches(0.8),
         [{"text": "在签名之前，先看清一笔交易到底会做什么", "size": 18,
           "color": MUTED, "align": PP_ALIGN.CENTER}])
    card(s, Inches(3.4), Inches(5.55), Inches(6.5), Inches(0.7), fill=BG_CARD)
    text(s, Inches(3.4), Inches(5.62), Inches(6.5), Inches(0.6),
         [{"text": "LXDAO × Monad 黑客松  ·  2026", "size": 15,
           "color": TEXT, "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)


# ================= 第 2 页：痛点 =================
def page_pain():
    s = slide_new(); bg(s)
    header(s, "问题背景", "链上交易的「盲签」之痛")
    panel(s, Inches(0.6), Inches(1.95), Inches(7.3), Inches(4.6),
          "常见的 5 种风险",
          bullets([
              "授权无感知：点一下「授权」，可能交出无限额度",
              "收款方被掉包：看似转给朋友，实际转给黑客地址",
              "金额被膨胀：1 MON 变 100 MON，签名时已晚",
              "夹带授权：一笔操作里偷偷塞进授权调用",
              "签名即失控：交易上链不可逆，错了追不回",
          ], size=16, gap=12), accent=DANGER, title_size=20)
    card(s, Inches(8.15), Inches(1.95), Inches(4.55), Inches(4.6),
         fill=BG_CARD, line=AMBER, radius=0.1)
    accent_bar(s, Inches(8.15), Inches(1.95), Inches(4.6), AMBER)
    text(s, Inches(8.45), Inches(2.2), Inches(4.0), Inches(0.6),
         [{"text": "真正的后果", "size": 20, "color": AMBER, "bold": True}])
    text(s, Inches(8.45), Inches(3.0), Inches(4.0), Inches(1.6),
         [{"text": "0", "size": 80, "color": DANGER, "bold": True,
           "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(8.45), Inches(4.7), Inches(4.0), Inches(0.5),
         [{"text": "后悔药", "size": 18, "color": MUTED,
           "align": PP_ALIGN.CENTER}])
    text(s, Inches(8.45), Inches(5.4), Inches(4.0), Inches(1.0),
         [{"text": "用户往往在签名之后，才发现资产已经被转走。",
           "size": 15, "color": TEXT, "align": PP_ALIGN.CENTER,
           "line_spacing": 1.2}])
    footer(s, 2)


# ================= 第 3 页：定位 =================
def page_position():
    s = slide_new(); bg(s)
    header(s, "产品定位", "MonadLens 是什么")
    text(s, Inches(0.6), Inches(1.75), Inches(12.1), Inches(0.8),
         [{"text": "一个把「签名前后果」可视化、可对话、可拦截的链上安全助手。",
           "size": 19, "color": TEXT}], anchor=MSO_ANCHOR.MIDDLE)
    cols = [
        ("📊", "实时看板", "把链上数据变成看得懂的画面", CYAN),
        ("🤖", "AI 对话", "用自然语言规划交易、演示攻击", PURPLE),
        ("🛡️", "后果透镜", "签名前模拟并拦截危险操作", AMBER),
    ]
    w = Inches(3.85); gap = Inches(0.27); left = Inches(0.6); top = Inches(2.75)
    for i, (icon, title, desc, acc) in enumerate(cols):
        l = left + i * (w + gap)
        card(s, l, top, w, Inches(3.2), fill=BG_CARD)
        accent_bar(s, l, top, Inches(3.2), acc)
        text(s, l, top + Inches(0.35), w, Inches(0.9),
             [{"text": icon, "size": 40, "color": acc, "align": PP_ALIGN.CENTER}],
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, l, top + Inches(1.35), w, Inches(0.6),
             [{"text": title, "size": 22, "color": WHITE, "bold": True,
               "align": PP_ALIGN.CENTER}])
        text(s, l + Inches(0.2), top + Inches(2.05), w - Inches(0.4), Inches(1.0),
             [{"text": desc, "size": 15, "color": MUTED, "align": PP_ALIGN.CENTER,
               "line_spacing": 1.2}])
    footer(s, 3)


# ================= 第 4 页：架构 =================
def page_arch():
    s = slide_new(); bg(s)
    header(s, "整体架构", "三栏一体化：数据 · 对话 · 防护")
    cols = [
        ("左栏", "实时看板", "Monad 链上实时数据\nTPS / 出块 / Gas / 活跃榜", CYAN),
        ("中栏", "Agent 双脑", "自然语言交互\n规则引擎 + DeepSeek LLM", PURPLE),
        ("右栏", "后果透镜", "签名前 Moss 模拟\n预览资金流向与授权", AMBER),
    ]
    w = Inches(3.85); gap = Inches(0.27); left = Inches(0.6); top = Inches(2.0)
    for i, (tag, title, desc, acc) in enumerate(cols):
        l = left + i * (w + gap)
        card(s, l, top, w, Inches(3.1), fill=BG_CARD)
        accent_bar(s, l, top, Inches(3.1), acc)
        text(s, l + Inches(0.2), top + Inches(0.2), w - Inches(0.4), Inches(0.4),
             [{"text": tag, "size": 13, "color": acc, "bold": True}])
        text(s, l, top + Inches(0.7), w, Inches(0.7),
             [{"text": title, "size": 24, "color": WHITE, "bold": True,
               "align": PP_ALIGN.CENTER}])
        text(s, l + Inches(0.25), top + Inches(1.6), w - Inches(0.5), Inches(1.3),
             [{"text": desc, "size": 15, "color": MUTED, "align": PP_ALIGN.CENTER,
               "line_spacing": 1.3}])
    card(s, Inches(0.6), Inches(5.45), Inches(12.1), Inches(1.05), fill=BG_CARD,
         line=PURPLE, radius=0.1)
    text(s, Inches(0.85), Inches(5.55), Inches(11.6), Inches(0.9),
         [{"text": "数据流：钱包 / 链上  →  看板可视化  →  Agent 规划  →  Moss 模拟  →  "
                  "透镜拦截 / 放行  →  签名广播  →  链上验证",
           "size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "line_spacing": 1.3}],
         anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 4)


# ================= 第 5 页：功能① 看板 =================
def page_feat_dashboard():
    s = slide_new(); bg(s)
    header(s, "核心功能 ①", "实时链上看板")
    panel(s, Inches(0.6), Inches(1.95), Inches(7.0), Inches(4.55),
          "你能看到什么",
          bullets([
              "TPS 实时估算（约 40 tx/s）",
              "出块间隔约 0.3s，毫秒级刷新",
              "Gas 费用趋势图，提前感知拥堵",
              "活跃地址 / 合约榜 Top 5 实时更新",
              "WS 与 HTTP 竞速，慢节点也不卡死",
          ], size=16, gap=11), accent=CYAN, title_size=20)
    card(s, Inches(7.85), Inches(1.95), Inches(4.85), Inches(4.55), fill=BG_CARD)
    accent_bar(s, Inches(7.85), Inches(1.95), Inches(4.55), CYAN)
    text(s, Inches(8.1), Inches(2.1), Inches(4.3), Inches(0.5),
         [{"text": "Gas 趋势（示意）", "size": 17, "color": WHITE, "bold": True}])
    for gy in range(4):
        y = Inches(2.8 + gy * 0.85)
        ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.15), y,
                                Inches(4.4), Inches(0.012))
        ln.fill.solid(); ln.fill.fore_color.rgb = RGBColor(0x33, 0x2E, 0x55)
        ln.line.fill.background()
    line_pts = [(8.2, 5.6), (9.0, 4.7), (9.8, 5.1), (10.6, 3.9),
                (11.4, 4.4), (12.2, 3.4)]
    prev = None
    for (x, y) in line_pts:
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.06), Inches(y - 0.06),
                                 Inches(0.12), Inches(0.12))
        dot.fill.solid(); dot.fill.fore_color.rgb = CYAN; dot.line.fill.background()
        if prev:
            px, py = prev
            conn = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(px), Inches(py),
                                      Inches(0.05), Inches(0.05))
            conn.fill.solid(); conn.fill.fore_color.rgb = CYAN
            conn.line.fill.background()
        prev = (x, y)
    text(s, Inches(8.1), Inches(6.05), Inches(4.3), Inches(0.4),
         [{"text": "实时数据，一眼看懂网络状态", "size": 12, "color": MUTED,
           "align": PP_ALIGN.CENTER}])
    footer(s, 5)


# ================= 第 6 页：功能② Agent =================
def page_feat_agent():
    s = slide_new(); bg(s)
    header(s, "核心功能 ②", "Agent 双脑 AI 对话")
    panel(s, Inches(0.6), Inches(1.95), Inches(5.9), Inches(2.55),
          "🧠 规则引擎", bullets([
              "快、确定，覆盖已知意图",
              "交易规划、攻击演示一键触发",
          ], size=15, gap=9), accent=CYAN, title_size=19)
    panel(s, Inches(6.8), Inches(1.95), Inches(5.9), Inches(2.55),
          "✨ DeepSeek LLM", bullets([
              "自然、智能，理解自由提问",
              "解释「为什么先模拟再签名」等",
          ], size=15, gap=9), accent=PURPLE, title_size=19)
    card(s, Inches(0.6), Inches(4.75), Inches(12.1), Inches(1.75), fill=BG_CARD,
         line=PURPLE, radius=0.1)
    accent_bar(s, Inches(0.6), Inches(4.75), Inches(1.75), PURPLE)
    text(s, Inches(0.9), Inches(4.9), Inches(11.5), Inches(0.5),
         [{"text": "🔄 自动降级：LLM 不可用时切回规则引擎，永不掉线",
           "size": 17, "color": WHITE, "bold": True}])
    text(s, Inches(0.9), Inches(5.5), Inches(11.5), Inches(0.9),
         [{"text": "两条大脑互为备份，演示现场再也不怕「AI 抽风」导致功能失灵。",
           "size": 15, "color": MUTED, "line_spacing": 1.2}])
    footer(s, 6)


# ================= 第 7 页：功能③ 透镜 =================
def page_feat_lens():
    s = slide_new(); bg(s)
    header(s, "核心功能 ③", "后果透镜（签名前模拟）")
    panel(s, Inches(0.6), Inches(1.95), Inches(6.0), Inches(4.55),
          "它模拟什么", bullets([
              "资金流向：谁转出、谁收到、多少",
              "授权范围：交出的是定额还是无限",
              "余额变化：签名前后的净变动",
              "收款方对账：有没有偷偷换地址",
          ], size=16, gap=12), accent=AMBER, title_size=20)
    panel(s, Inches(6.85), Inches(1.95), Inches(5.85), Inches(4.55),
          "为什么重要", bullets([
              "基于 Moss SDK 的 debug_traceCall 真模拟",
              "主网 0 MON 也能跑（内部预充 1000 MON）",
              "把「看不见的后果」变成「看得见的清单」",
              "先把风险看清楚，再决定要不要签",
          ], size=16, gap=12), accent=CYAN, title_size=20)
    footer(s, 7)


# ================= 第 8 页：功能④ 攻击护栏 =================
def page_feat_guard():
    s = slide_new(); bg(s)
    header(s, "核心功能 ④", "5 类攻击护栏 — 签名前拦截")
    items = [
        ("🔄", "收款方掉包"), ("🔓", "夹带授权"), ("💥", "金额膨胀"),
        ("♾️", "无限授权"), ("✍️", "篡改后签名"),
    ]
    w = Inches(2.28); gap = Inches(0.17); left = Inches(0.6); top = Inches(2.0)
    for i, (icon, name) in enumerate(items):
        l = left + i * (w + gap)
        card(s, l, top, w, Inches(1.85), fill=BG_CARD)
        text(s, l, top + Inches(0.2), w, Inches(0.8),
             [{"text": icon, "size": 30, "color": AMBER, "align": PP_ALIGN.CENTER}],
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, l, top + Inches(0.95), w, Inches(0.5),
             [{"text": name, "size": 15, "color": WHITE, "bold": True,
               "align": PP_ALIGN.CENTER}])
        text(s, l, top + Inches(1.4), w, Inches(0.4),
             [{"text": "全部 blocked", "size": 12, "color": CYAN,
               "align": PP_ALIGN.CENTER, "bold": True}])
    card(s, Inches(0.6), Inches(4.25), Inches(12.1), Inches(2.25), fill=BG_CARD,
         line=DANGER, radius=0.1)
    accent_bar(s, Inches(0.6), Inches(4.25), Inches(2.25), DANGER)
    text(s, Inches(0.9), Inches(4.4), Inches(11.5), Inches(0.5),
         [{"text": "危险操作对比", "size": 18, "color": WHITE, "bold": True}])
    text(s, Inches(0.9), Inches(5.0), Inches(5.5), Inches(1.3),
         [{"text": "普通钱包", "size": 15, "color": MUTED, "bold": True},
          {"text": "直接放行 → 资产悄悄流失", "size": 15, "color": TEXT,
           "space_before": 6}])
    text(s, Inches(6.8), Inches(5.0), Inches(5.5), Inches(1.3),
         [{"text": "MonadLens", "size": 15, "color": CYAN, "bold": True},
          {"text": "标红告警 + 锁死签名按钮", "size": 15, "color": TEXT,
           "space_before": 6}])
    footer(s, 8)


# ================= 第 9 页：闭环 =================
def page_loop():
    s = slide_new(); bg(s)
    header(s, "端到端演示", "一条完整的可信交易闭环")
    steps = ["① 连钱包", "② Agent 对话", "③ Moss 模拟预览",
             "④ 用户确认", "⑤ 签名广播", "⑥ 链上验证"]
    w = Inches(1.92); gap = Inches(0.12); left = Inches(0.6); top = Inches(2.2)
    for i, st in enumerate(steps):
        l = left + i * (w + gap)
        card(s, l, top, w, Inches(1.4), fill=BG_CARD)
        text(s, l, top + Inches(0.1), w, Inches(1.2),
             [{"text": st, "size": 15, "color": WHITE, "bold": True,
               "align": PP_ALIGN.CENTER, "line_spacing": 1.15}],
             anchor=MSO_ANCHOR.MIDDLE)
        if i < len(steps) - 1:
            ar = s.shapes.add_shape(MSO_SHAPE.CHEVRON, l + w - Inches(0.02),
                                    top + Inches(0.45), Inches(0.16), Inches(0.5))
            ar.fill.solid(); ar.fill.fore_color.rgb = PURPLE
            ar.line.fill.background()
    card(s, Inches(0.6), Inches(4.1), Inches(12.1), Inches(2.2), fill=BG_CARD,
         line=CYAN, radius=0.1)
    accent_bar(s, Inches(0.6), Inches(4.1), Inches(2.2), CYAN)
    text(s, Inches(0.9), Inches(4.3), Inches(11.5), Inches(0.6),
         [{"text": "✅ 已实测：在 Monad 测试网真实签名并上链成功",
           "size": 18, "color": CYAN, "bold": True}])
    text(s, Inches(0.9), Inches(5.0), Inches(11.5), Inches(1.1),
         [{"text": "交易哈希已在 MonadScan 验证通过 —— 模拟、签名、广播、验证全链路跑通，"
                   "不是 Demo 截图，是真上链。",
           "size": 15, "color": MUTED, "line_spacing": 1.3}])
    footer(s, 9)


# ================= 第 10 页：技术栈 =================
def page_stack():
    s = slide_new(); bg(s)
    header(s, "技术栈", "现代、稳健、可演示")
    groups = [
        ("前端 / 框架", ["Next.js 16", "React 19", "Tailwind v4", "TypeScript"], PURPLE),
        ("链上 / 钱包", ["wagmi v3", "viem v2", "zustand", "Moss SDK"], CYAN),
        ("AI / 可视化", ["DeepSeek LLM", "规则引擎", "recharts", "debug_traceCall"], AMBER),
    ]
    w = Inches(3.85); gap = Inches(0.27); left = Inches(0.6); top = Inches(2.0)
    for i, (title, items, acc) in enumerate(groups):
        l = left + i * (w + gap)
        panel(s, l, top, w, Inches(3.6), title, bullets(items, size=17, gap=12),
              accent=acc, title_size=19, body_size=17)
    footer(s, 10)


# ================= 第 11 页：演示脚本 =================
def page_demo():
    s = slide_new(); bg(s)
    header(s, "如何体验", "3 步跑通演示（测试网 · 免费）")
    steps = [
        ("1", "连钱包 + 一键加网", "连接 MetaMask/OKX，点「添加网络」自动灌入 Monad 测试网。", CYAN),
        ("2", "领免费测试币", "打开 faucet.monad.xyz，领取免费测试 MON（无真实价值）。", PURPLE),
        ("3", "对话 → 模拟 → 签名", "点建议按钮触发模拟，看透镜预览，签名后到 MonadScan 验证。", AMBER),
    ]
    top = Inches(2.0); h = Inches(1.45)
    for i, (num, title, desc, acc) in enumerate(steps):
        t = top + i * (h + Inches(0.2))
        card(s, Inches(0.6), t, Inches(12.1), h, fill=BG_CARD)
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.85), t + Inches(0.32),
                               Inches(0.8), Inches(0.8))
        c.fill.solid(); c.fill.fore_color.rgb = acc; c.line.fill.background()
        text(s, Inches(0.85), t + Inches(0.32), Inches(0.8), Inches(0.8),
             [{"text": num, "size": 30, "color": WHITE, "bold": True,
               "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(1.95), t + Inches(0.18), Inches(10.4), Inches(0.5),
             [{"text": title, "size": 19, "color": WHITE, "bold": True}])
        text(s, Inches(1.95), t + Inches(0.72), Inches(10.4), Inches(0.6),
             [{"text": desc, "size": 14, "color": MUTED}])
    card(s, Inches(0.6), Inches(6.35), Inches(12.1), Inches(0.55), fill=BG_CARD,
         line=CYAN, radius=0.1)
    text(s, Inches(0.6), Inches(6.35), Inches(12.1), Inches(0.55),
         [{"text": "提示：测试网代币免费、无真实价值，可放心演示每一步。",
           "size": 13, "color": CYAN, "align": PP_ALIGN.CENTER}],
         anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 11)


# ================= 第 12 页：结尾 =================
def page_end():
    s = slide_new(); bg(s)
    card(s, 0, 0, SW, Inches(0.2), fill=PURPLE, radius=0)
    text(s, Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.2),
         [{"text": "看清后果，才能放心签名", "size": 44, "color": WHITE,
           "bold": True, "align": PP_ALIGN.CENTER}])
    text(s, Inches(0.9), Inches(3.8), Inches(11.5), Inches(0.7),
         [{"text": "MonadLens — 让你的每一笔签名都心里有数",
           "size": 20, "color": PURPLE, "align": PP_ALIGN.CENTER}])
    text(s, Inches(0.9), Inches(5.4), Inches(11.5), Inches(0.5),
         [{"text": "致谢  LXDAO × Monad  黑客松", "size": 15, "color": MUTED,
           "align": PP_ALIGN.CENTER}])


# ---------------- 组装 ----------------
page_cover(); page_pain(); page_position(); page_arch()
page_feat_dashboard(); page_feat_agent(); page_feat_lens()
page_feat_guard(); page_loop(); page_stack(); page_demo(); page_end()

OUT = r"C:\Users\EmptyTouch\WorkBuddy\2026-08-07-21-45-54\monadlens\MonadLens演示.pptx"
prs.save(OUT)
print("SAVED", OUT, "slides=", len(prs.slides._sldIdLst))
