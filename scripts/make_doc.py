# -*- coding: utf-8 -*-
"""Generate MonadLens project submission document in HTML + DOCX (rich-text friendly)."""
from __future__ import annotations
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT_DIR = r"C:\Users\EmptyTouch\WorkBuddy\2026-08-07-21-45-54\monadlens"

SECTIONS = [
    {
        "heading": "一、项目定位",
        "body": [
            "MonadLens（中文名「透视链」）是面向 Monad 生态的对话式链上助手 + 签名前后果透镜。",
            (
                "用户用中文描述想做的事，AI Agent 负责构造交易；"
                "但在用户签名之前，系统通过 Moss SDK 在真实链上状态中"
                "模拟执行这笔交易，把「转出什么、收到什么、授权给谁、"
                "钱最终落到哪个地址」用人话摊开，并对收款方掉包、"
                "金额膨胀、无限授权等攻击实时拦截。"
            ),
            "一句话定位：让用户在签名之前，真正看懂这笔交易会带来的后果。",
        ],
    },
    {
        "heading": "二、主要亮点",
        "body": [
            ("签名前后果透镜（核心能力）", True),
            (
                "基于 Moss 的 debug_traceCall 在真实链上状态模拟，"
                "把资金流出 / 流入、授权明细、收款方对账、安全告警"
                "逐项展示，签名前即知后果。"
            ),
            ("应用层「收款方对账」（差异化能力）", True),
            (
                "Moss 的信封约束只保证「转出多少」，不约束「转给谁」。"
                "MonadLens 补上这一层：当界面告诉用户的收款方，与实际 "
                "calldata 的收款方不一致时，产生 UNDECLARED_RECIPIENT 告警并拦截——"
                "抓住「金额一分不差、收款地址被换」的隐蔽攻击。"
            ),
            ("五类攻击护栏，全部拦截", True),
            (
                "收款方掉包、无限授权、夹带授权、金额膨胀、封印后篡改。"
                "任一命中即判定 blocked，签名按钮直接锁死——不是事后报警，是事前挡住。"
            ),
            ("Agent 双脑架构（规则引擎 + LLM）", True),
            (
                "内置规则引擎覆盖常见操作（转账 / Wrap / Swap / 授权），"
                "无需 API Key 即可运行；配置 DeepSeek 等 OpenAI 兼容接口后自动升级为 LLM 路由，"
                "LLM 失败时无缝降级回规则引擎，功能不中断。"
            ),
            ("实时链上看板", True),
            (
                "左侧栏实时展示 Monad 网络 TPS、出块间隔、Gas 费趋势、"
                "活跃地址/合约 Top5，WebSocket 与 HTTP 双通道竞速，数据来自真实主网/测试网。"
            ),
            ("一键添加 Monad 网络到钱包", True),
            (
                "用户无需手动配 RPC —— 点一下即可通过 wallet_addEthereumChain "
                "把 Monad 主网或测试网络灌入 MetaMask / OKX 等钱包，并自动切换。"
            ),
            ("测试网免费完整闭环", True),
            (
                "切换到测试网模式后，用户可领免费水龙头币，走完「连钱包 → 对话 → "
                "模拟预览 → 签名广播 → 区块链浏览器验证」的完整流程，零成本体验全部功能。"
            ),
        ],
    },
    {
        "heading": "三、使用场景",
        "body": [
            ("防骗场景", True),
            (
                "当有人发来「请向此地址转账 X MON」的请求时，"
                "用户在 MonadLens 中发起操作，透镜会在签名前揭示真实收款方与预期是否一致、"
                "是否有隐藏授权等风险。"
            ),
            ("安全演示与教育", True),
            (
                "讲师/开发者用演示模式展示五类攻击手法及系统如何拦截，"
                "帮助受众理解链上交易的风险面。"
            ),
            ("开发者接入", True),
            "开源项目，开发者可基于 Moss SDK + 自身规则扩展更多检测维度。",
            ("零成本体验", True),
            "测试网模式下无需持有真币即可完整体验从对话到上链的全流程。",
        ],
    },
    {
        "heading": "四、技术信息",
        "body": [
            ("演示地址", True), "部署后获得（支持 Vercel 一键导入 GitHub 仓库）",
            ("GitHub", True), "https://github.com/<你的用户名>/monadlens （待 push）",
            None,
            ("技术栈", True),
            "前端：Next.js 16 + React 19 + Tailwind CSS v4 + TypeScript + recharts v3",
            "链上交互：wagmi v3 + viem v2 + Moss Onchain Agent SDK",
            "AI：规则引擎 + DeepSeek / OpenAI 兼容 LLM（可选）",
            "实时数据：WebSocket + HTTP RPC 多端点竞速 fallback",
            None,
            ("运行命令", True),
            "git clone <仓库地址> && cd monadlens",
            "npm install",
            "cp .env.example .env.local   # 按需填写（LLM 可不配）",
            "npm run dev                  # http://localhost:3000",
            None,
            ("网络切换", True),
            "修改 .env.local 中 NEXT_PUBLIC_MONAD_NETWORK=testnet|mainnet 后重新 npm run build 即可切换。默认 main网。",
            None,
        ],
    },
]

SHORT_DESC = (
    "MonadLens（中文名「透视链」）是面向 Monad 生态的对话式链上安全助手。"
    "用户用中文描述交易意图，AI Agent 构造交易；但在签名前，"
    "系统通过 Moss SDK 在真实链上状态模拟执行，把资金流向、"
    "授权对象、收款方身份用人话展开，并对五类常见攻击（收款方掉包、"
    "无限授权、金额膨胀等）实时拦截。"
    "支持 Agent 双脑架构（规则引擎 + DeepSeek LLM 自动降级）、"
    "实时链上看板、一键添加 Monad 网络到钱包。"
    "测试网可免费跑通完整闭环（对话→模拟→签名→上链→验证）。"
    "技术栈：Next.js 16 / React 19 / wagmi v3 / Moss SDK / DeepSeek。"
)


def gen_html() -> str:
    lines = [
        "<!DOCTYPE html><html lang=zh-CN><head><meta charset=UTF-8>",
        "<title>MonadLens · 透视链 项目文档</title>",
        "<style>",
        "body{font-family:'Microsoft YaHei','PingFang SC',sans-serif;max-width:780px;margin:36px auto;padding:0 20px;color:#1a1a1a;line-height:1.85;font-size:15px;}",
        "h2{color:#6B21A8;border-bottom:2px solid #6B21A8;padding-bottom:6px;margin-top:34px;font-size:19px;}",
        "h3{color:#1e40af;margin-top:18px;font-size:15.5px;}",
        "p{margin:7px 0;text-indent:0;}",
        "ul,ol{padding-left:22px;margin:6px 0;}",
        "li{margin:4px 0;}",
        "strong{color:#111;}",
        "code{background:#f3f4f6;padding:1px 6px;border-radius:3px;font-size:13px;}",
        ".box{background:#eff6ff;border-left:4px solid #2563eb;padding:14px 18px;margin:22px 0;border-radius:0 6px 6px 0;font-size:14.5px;line-height:1.9;}",
        "</style></head><body>",
    ]
    for sec in SECTIONS:
        lines.append(f"<h2>{sec['heading']}</h2>")
        for item in sec["body"]:
            if item is None:
                lines.append("<br>")
            elif isinstance(item, tuple):
                txt, bold = item
                tag = "h3" if bold else "p"
                lines.append(f"<{tag}>{txt}</{tag}>")
            else:
                lines.append(f"<p>{item}</p>")

    lines.append(f'<div class="box"><strong>报名简介（可直接复制）：</strong><br><br>{SHORT_DESC}</div>')
    lines.append("</body></html>")
    return "\n".join(lines)


def _rf(run, size=11, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
        "\u5fae\u8f6f\u96c5\u9ed1",
    )
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def gen_docx() -> Document:
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)

    # title
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _rf(t.add_run("MonadLens \u00b7 \u900f\u89c6\u94fe \u2014 \u9879\u76ee\u6587\u6863"), size=22, bold=True, color=(107, 33, 168))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _rf(sub.add_run("LXDAO \u00d7 Monad \u9ed1\u5ba2\u677e\u53c2\u8d5b\u4f5c\u54c1"), size=12, color=(107, 114, 128))
    doc.add_paragraph()

    for sec in SECTIONS:
        h = doc.add_paragraph()
        _rf(h.add_run(sec["heading"]), size=16, bold=True, color=(107, 33, 168))
        for item in sec["body"]:
            if item is None:
                doc.add_paragraph()
            elif isinstance(item, tuple):
                txt, bold = item
                p = doc.add_paragraph()
                _rf(p.add_run(txt), size=12 if bold else 11, bold=bold, color=(30, 64, 175) if bold else None)
                p.paragraph_format.space_before = Pt(5)
            else:
                p = doc.add_paragraph()
                _rf(p.add_run(item), size=11)
                p.paragraph_format.space_after = Pt(2)

    doc.add_paragraph()
    bp = doc.add_paragraph()
    _rf(bp.add_run("\u62a5\u540d\u7b80\u4ecb\uff08\u53ef\u76f4\u63a5\u590d\u5236\uff09\uff1a"), size=12, bold=True)
    dp = doc.add_paragraph()
    _rf(dp.add_run(SHORT_DESC), size=11)

    return doc


if __name__ == "__main__":
    html_path = os.path.join(OUT_DIR, "\u9879\u76ee\u62a5\u540d\u6587\u6863.html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(gen_html())
    print(f"HTML -> {html_path}")

    docx_path = os.path.join(OUT_DIR, "\u9879\u76ee\u62a5\u540d\u6587\u6863.docx")
    gen_docx().save(docx_path)
    print(f"DOCX -> {docx_path}")
